"""Ask Poolside answer → .docx in the briefing's editorial grammar
(pipeline/briefing.py): scope as the masthead, an italic kicker with the
date, the question as the headline, eyebrow section labels over accent
hairlines, alignment-group titles in accent, superscript `[n]` citations
that resolve in a numbered Sources list at the end."""
from __future__ import annotations

import io
import re
from datetime import date, datetime

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import pipeline.brand as brand
import pipeline.db as db
from pipeline.briefing import (
    _CYAN,
    _CYAN_HEX,
    _GRAY_MID_HEX,
    _GRAY_TEXT,
    _INK,
    _INK_SOFT,
    _LABEL,
    _eyebrow,
    _unescape_inline,
    _v2_link,
    _v2_page_number,
    _v2_pborder,
    _v2_right_tab,
    _v2_run,
    _v2_spacing,
)
from pipeline.ferc_client import docinfo_url

_DETAIL_LABEL = {"brief": "Brief answer", "standard": "Standard answer",
                 "deep": "Deep memo"}

# Inline grammar of a stored answer: bold/italic toggles, `[n]` citation
# runs (possibly several back to back), page references after a citation.
_TOKEN_RE = re.compile(r"(\*\*|(?<!\*)\*(?!\*)|(?:\[\d+\])+|\(pp?\.\s*[\d,\s–-]+\))")
_MD_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")
_WHOLE_BOLD_RE = re.compile(r"^\*\*(.+?)\*\*:?\s*$")


def _masthead_and_kicker(scope: dict | None) -> tuple[str, str]:
    """(masthead, kicker): the scope as the big cover title, the docket
    title (when there is one) as the italic line under it."""
    scope = scope or {}
    dockets = scope.get("dockets") or []
    if dockets:
        nums = ", ".join(d.get("docket_number", "?") for d in dockets)
        masthead = f"FERC Docket{'s' if len(dockets) > 1 else ''} {nums}"
        titles = [d.get("title") for d in dockets if d.get("title")]
        kicker = titles[0] if len(titles) == 1 else "Research memo"
        return masthead, kicker
    if scope.get("corpus") == "dockets":
        return "Tracked FERC Dockets", "Research memo"
    if scope.get("corpus") == "meetings":
        return "Stakeholder Meetings", "Research memo"
    return "Meetings & FERC Dockets", "Research memo"


def _scope_footer(scope: dict | None) -> str:
    scope = scope or {}
    dockets = scope.get("dockets") or []
    if dockets:
        return ", ".join(d.get("docket_number", "?") for d in dockets)
    return {"dockets": "Tracked dockets", "meetings": "Meetings"}.get(
        scope.get("corpus") or "", "Meetings + dockets")


def _inline(para, text: str, *, size=brand.SZ_BODY, color=_INK_SOFT,
            bold: bool = False, italic: bool = False) -> None:
    """Runs for one line: **bold** / *italic* toggles, `[n]` as small accent
    superscripts (what a memo's reader expects of a citation), page refs
    small and grey."""
    text = _MD_LINK_RE.sub(r"\1", text)
    is_bold, is_italic = bold, italic
    for part in _TOKEN_RE.split(text):
        if not part:
            continue
        if part == "**":
            is_bold = not is_bold
            continue
        if part == "*":
            is_italic = not is_italic
            continue
        if part.startswith("[") and part.endswith("]"):
            nums = re.findall(r"\d+", part)
            r = para.add_run(",".join(nums))
            r.font.name = brand.BODY_FONT
            r.font.size = Pt(8)
            r.font.superscript = True
            r.font.color.rgb = _CYAN
            continue
        if part.startswith("(p"):
            _v2_run(para, " " + part.strip(), size=brand.SZ_LINK, color=_GRAY_TEXT)
            continue
        _v2_run(para, _unescape_inline(part), size=size, bold=is_bold,
                italic=is_italic, color=color)


def _body(doc: Document, answer_md: str) -> None:
    """The answer in the briefing's section grammar.

    `##`/`###` → eyebrow label over a hairline (BOTTOM LINE, BACKGROUND …);
    a whole-line bold paragraph (the deep memo's alignment groups,
    "**A. Seek express grandfathering…**") → accent section title;
    bullets → accent dash with hanging indent; the first paragraph under
    an eyebrow sits in ink, the rest in soft ink — as the briefing does."""
    first_in_section = True
    for raw in (answer_md or "").splitlines():
        s = raw.strip()
        if not s or re.fullmatch(r"-{3,}", s):
            continue
        m = re.match(r"^#{1,4}\s+(.+?)\s*$", s)
        if m:
            _eyebrow(doc, m.group(1).strip().upper())
            first_in_section = True
            continue
        m = _WHOLE_BOLD_RE.match(s)
        if m and "[" not in m.group(1)[:1]:
            p = doc.add_paragraph()
            _v2_spacing(p, before=Pt(14), after=Pt(5), line=brand.LINE_SPACING)
            p.paragraph_format.keep_with_next = True
            _inline(p, m.group(1), size=brand.SZ_GROUP, color=_CYAN, bold=True)
            first_in_section = True
            continue
        if re.match(r"^[-*•]\s+", s):
            p = doc.add_paragraph()
            _v2_spacing(p, before=Pt(0), after=Pt(3), line=brand.LINE_SPACING)
            p.paragraph_format.left_indent = Pt(14)
            p.paragraph_format.first_line_indent = Pt(-14)
            p.paragraph_format.widow_control = True
            _v2_run(p, "–  ", size=brand.SZ_BODY, bold=True, color=_CYAN)
            _inline(p, re.sub(r"^[-*•]\s+", "", s), color=_INK_SOFT)
            first_in_section = False
            continue
        if s.startswith(">"):
            p = doc.add_paragraph()
            _v2_spacing(p, before=Pt(4), after=Pt(8), line=brand.LINE_SPACING)
            _v2_pborder(p, "left", 12, _CYAN_HEX, space=8)
            p.paragraph_format.left_indent = Pt(10)
            _inline(p, re.sub(r"^>\s?", "", s), color=_INK_SOFT, italic=True)
            first_in_section = False
            continue
        p = doc.add_paragraph()
        _v2_spacing(p, before=Pt(0), after=Pt(7), line=brand.LINE_SPACING)
        _inline(p, s, color=_INK if first_in_section else _INK_SOFT)
        first_in_section = False


def _source_line(src: dict) -> tuple[str, str | None]:
    """(label, url) for one serialized source."""
    et = src.get("entity_type")
    if src.get("docket_number"):
        head = f"FERC docket {src['docket_number']}"
        if et == "docket":
            return f"{head} — state of play", None
        bits = [src.get("accession_number") or "filing"]
        if src.get("document_class"):
            bits.append(src["document_class"])
        if src.get("filed_date"):
            bits.append(f"filed {src['filed_date']}")
        if src.get("authors"):
            bits.append(f"by {src['authors']}")
        if src.get("tier") == "document" and src.get("filename"):
            pages = src.get("pages") or []
            pg = f", p. {', '.join(str(p) for p in pages)}" if pages else ""
            bits.append(f"excerpt: {src['filename']}{pg}")
        elif src.get("description"):
            bits.append(str(src["description"])[:140])
        url = docinfo_url(src["accession_number"]) if src.get("accession_number") else None
        return f"{head} — " + " · ".join(bits), url
    head = f"{src.get('type_short') or '?'} meeting {src.get('meeting_date') or ''}".strip()
    if src.get("tier") == "document":
        item = f"agenda {src['item_id']} · " if src.get("item_id") else ""
        return f"{head} — {item}excerpt: {src.get('filename') or 'document'}", None
    if src.get("item_id"):
        return f"{head} — agenda {src['item_id']}: {src.get('item_title') or 'Untitled item'}", None
    return f"{head} — meeting briefing", None


def generate_ask_docx_bytes(ask_id: int) -> tuple[bytes, str]:
    """Render one ask_log exchange; returns (bytes, suggested_filename).
    Raises ValueError when the exchange is missing."""
    row = db.get_ask_log(ask_id)
    if not row:
        raise ValueError(f"Ask #{ask_id} not found")
    question = (row.get("question") or "").strip()
    answer_md = (row.get("answer_md") or "").strip()
    sources = row.get("sources") or []
    scope = row.get("scope") or {}
    created = row.get("created_at")
    when = (created.strftime("%B %-d, %Y") if isinstance(created, (date, datetime))
            else date.today().strftime("%B %-d, %Y"))
    masthead, kicker = _masthead_and_kicker(scope)
    prov = [_DETAIL_LABEL.get(row.get("detail") or "", "")]
    if row.get("model_id"):
        prov.append(str(row["model_id"]))
    if row.get("effort"):
        prov.append(f"{row['effort']} effort")
    prov.append(f"{len(sources)} source{'s' if len(sources) != 1 else ''}")
    if scope.get("depth") == "documents":
        prov.append("summaries + documents")
    prov_txt = " · ".join(p for p in prov if p)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = brand.MARGIN_TOPBOT; sec.bottom_margin = brand.MARGIN_TOPBOT
    sec.left_margin = brand.MARGIN_SIDE; sec.right_margin = brand.MARGIN_SIDE
    content_w = (sec.page_width - sec.left_margin - sec.right_margin) // 635

    style = doc.styles["Normal"]
    style.font.name = brand.BODY_FONT; style.font.size = brand.SZ_BODY
    style.font.color.rgb = _INK_SOFT
    style.paragraph_format.line_spacing = brand.LINE_SPACING

    sec.different_first_page_header_footer = True
    for part in (sec.first_page_header, sec.first_page_footer, sec.header):
        if part.paragraphs:
            part.paragraphs[0].clear()
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    fp.clear()
    pPr = fp._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    for val, pos in (("center", content_w // 2), ("right", content_w)):
        t = OxmlElement("w:tab"); t.set(qn("w:val"), val); t.set(qn("w:pos"), str(pos))
        tabs.append(t)
    pPr.append(tabs)
    _v2_run(fp, "Research Memo", size=brand.SZ_FOOTER, color=_GRAY_TEXT, font=_LABEL)
    fp.add_run("\t")
    _v2_run(fp, "Page ", size=brand.SZ_FOOTER, color=_GRAY_TEXT, font=_LABEL)
    pr = fp.add_run(); pr.font.name = _LABEL; pr.font.size = brand.SZ_FOOTER
    pr.font.color.rgb = _GRAY_TEXT
    _v2_page_number(pr)
    fp.add_run("\t")
    _v2_run(fp, _scope_footer(scope), size=brand.SZ_FOOTER, color=_GRAY_TEXT, font=_LABEL)
    _v2_pborder(fp, "top", 6, _CYAN_HEX, space=4)

    if doc.paragraphs:
        doc.paragraphs[0]._p.getparent().remove(doc.paragraphs[0]._p)

    # ── Cover: the briefing's header, with the scope as the masthead ────
    p = doc.add_paragraph(); _v2_spacing(p, before=Pt(0), after=Pt(0))
    _v2_pborder(p, "top", 30, _CYAN_HEX)
    p = doc.add_paragraph(); _v2_spacing(p, before=Pt(10), after=Pt(2))
    _v2_run(p, "POOLSIDE RESEARCH DESK", size=brand.SZ_LABEL, bold=True,
            color=_CYAN, font=_LABEL, track=40)
    p = doc.add_paragraph(); _v2_spacing(p, before=Pt(0), after=Pt(6))
    p.paragraph_format.keep_with_next = True
    _v2_run(p, masthead, size=brand.SZ_MASTHEAD, bold=True, color=_INK)
    p = doc.add_paragraph(); _v2_spacing(p, before=Pt(0), after=Pt(12))
    _v2_right_tab(p, pos=content_w)
    _v2_run(p, kicker, size=brand.SZ_HEADLINE, color=_CYAN, italic=True)
    p.add_run("\t")
    _v2_run(p, when, size=brand.SZ_HEADLINE, color=_GRAY_TEXT)
    p = doc.add_paragraph(); _v2_spacing(p, before=Pt(0), after=Pt(0))
    _v2_pborder(p, "bottom", 4, _GRAY_MID_HEX)

    # The question is the headline — the same italic line the briefing
    # uses for its tagline.
    p = doc.add_paragraph(); _v2_spacing(p, before=Pt(10), after=Pt(0))
    _v2_run(p, question, size=brand.SZ_HEADLINE, color=_INK, italic=True)
    p = doc.add_paragraph(); _v2_spacing(p, before=Pt(6), after=Pt(0))
    _v2_run(p, prov_txt, size=brand.SZ_LINK, color=_GRAY_TEXT, font=_LABEL)

    # ── Answer ──────────────────────────────────────────────────────────
    _body(doc, answer_md)

    # ── Sources ─────────────────────────────────────────────────────────
    if sources:
        _eyebrow(doc, "SOURCES")
        p = doc.add_paragraph(); _v2_spacing(p, before=Pt(0), after=Pt(6))
        _v2_run(p, "Superscript numbers in the text refer to these entries.",
                size=brand.SZ_LINK, color=_GRAY_TEXT)
        for src in sources:
            label, url = _source_line(src)
            p = doc.add_paragraph()
            _v2_spacing(p, before=Pt(0), after=Pt(2), line=brand.LINE_SPACING)
            p.paragraph_format.left_indent = Pt(22)
            p.paragraph_format.first_line_indent = Pt(-22)
            _v2_run(p, f"{src.get('n', '?')}   ", size=Pt(10), bold=True,
                    color=_CYAN, font=_LABEL)
            _v2_run(p, label, size=Pt(10), color=_INK_SOFT)
            if url:
                p.add_run("  ")
                _v2_link(p, url, "eLibrary", size=brand.SZ_LINK, color=_CYAN)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    slug = re.sub(r"[^A-Za-z0-9]+", "-", question)[:48].strip("-") or "answer"
    return buf.read(), f"Ask_{ask_id}_{slug}.docx"
